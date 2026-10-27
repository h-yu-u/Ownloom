import asyncio
import io
import json
import tempfile
import unittest
import uuid
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from mlx_studio_runtime.adapters.persistence.attachment_preparation_repository import (
    PreparationRepository,
)
from mlx_studio_runtime.adapters.persistence.sqlite_repositories import (
    SQLiteRepositories,
)
from mlx_studio_runtime.domain.errors import RuntimeFault
from mlx_studio_runtime.domain.models import ModelCapabilities, fake_text_manifest
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from runtime_fixtures import create_test_runtime

SOURCE = Path(__file__).parents[1] / "services/runtime/src"


def pdf_bytes(text=True, blank=True):
    writer = PdfWriter()
    if text:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 250 Td (Readable evidence) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if blank:
        writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class AttachmentPreparationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.service = create_test_runtime(
            SQLiteRepositories(Path(self.directory.name) / "app.db"), SOURCE
        )

    async def asyncTearDown(self):
        await self.service.close()
        self.directory.cleanup()

    async def upload(self, data, name, **options):
        async def chunks():
            for start in range(0, len(data), 19):
                yield data[start : start + 19]

        return await self.service.attachments.upload(chunks(), name, **options)

    async def ready(self, item):
        preparation = item["processing"]["preparationId"]
        task = self.service.attachments.tasks.get(preparation)
        if task:
            await task
        return self.service.attachments.metadata(item["id"])

    async def test_text_preparation_enters_text_model_context_and_persists_source(self):
        item = await self.ready(
            await self.upload("项目预算 120 万元".encode(), "预算.md")
        )
        self.assertEqual(item["processing"]["state"], "ready")
        conversation = await self.service.chat.create_conversation(
            "Test", "builtin-fake-text"
        )
        result = await self.service.chat.send_message(
            conversation["id"],
            "预算多少？",
            "builtin-fake-text",
            64,
            0,
            [item["id"]],
            visual_selection={"mode": "auto"},
        )
        await self.service.scheduler.tasks[result["jobId"]]
        job = self.service.queries.job(result["jobId"])
        self.assertEqual(job["state"], "succeeded")
        self.assertIn("120", job["text"])
        self.assertEqual(job["sources"][0]["attachmentId"], item["id"])
        self.assertEqual(
            job["sources"][0]["preparationId"], item["processing"]["preparationId"]
        )
        restored = self.service.queries.conversation(conversation["id"])
        self.assertEqual(restored["messages"][1]["sources"], job["sources"])
        self.assertNotIn("relativePath", json.dumps(restored["messages"][1]["sources"]))

    async def test_docx_body_and_table_are_in_order(self):
        document = Document()
        document.add_paragraph("Before")
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Table cell"
        document.add_paragraph("After")
        output = io.BytesIO()
        document.save(output)
        item = await self.ready(await self.upload(output.getvalue(), "report.docx"))
        preview = self.service.attachments.preview(
            item["id"], item["processing"]["preparationId"], "text"
        )
        self.assertEqual(
            [block["text"] for block in preview["blocks"]],
            ["Before", "Table cell | ", "After"],
        )
        self.assertTrue(all(block["page"] is None for block in preview["blocks"]))

    async def test_pdf_coverage_and_scanned_preview(self):
        item = await self.ready(await self.upload(pdf_bytes(), "mixed.pdf"))
        self.assertEqual(item["processing"]["state"], "partial")
        self.assertEqual(item["processing"]["coverage"]["unreadablePages"], [2])
        scanned = await self.ready(await self.upload(pdf_bytes(False), "scanned.pdf"))
        self.assertEqual(scanned["processing"]["state"], "needs_ocr")
        self.assertFalse(scanned["processing"]["retryable"])
        preview = self.service.attachments.preview(
            scanned["id"], scanned["processing"]["preparationId"], "text"
        )
        self.assertEqual(preview["blocks"], [])

    async def test_upload_retries_keep_id_and_conflicts_preserve_original(self):
        upload_id = str(uuid.uuid4())
        a, b = await asyncio.gather(
            self.upload(b"same", "a.txt", client_upload_id=upload_id),
            self.upload(b"same", "a.txt", client_upload_id=upload_id),
        )
        self.assertEqual(a["id"], b["id"])
        with self.assertRaises(RuntimeFault) as error:
            await self.upload(b"different", "a.txt", client_upload_id=upload_id)
        self.assertEqual(error.exception.info.code, "UPLOAD_IDEMPOTENCY_CONFLICT")
        self.assertEqual(
            self.service.attachments.read(a["id"])[0].read_bytes(), b"same"
        )

    async def test_gif_first_frame_and_corrupt_image_failure(self):
        output = io.BytesIO()
        Image.new("RGB", (20, 30), "red").save(
            output,
            "GIF",
            save_all=True,
            append_images=[Image.new("RGB", (20, 30), "blue")],
        )
        item = await self.ready(await self.upload(output.getvalue(), "animated.gif"))
        frozen = self.service.attachments.frozen(item["id"])
        with Image.open(self.service.attachments.image_path(frozen)) as image:
            self.assertEqual(image.getpixel((0, 0)), (255, 0, 0))
        self.assertIn("首帧", item["processing"]["warnings"][0])
        corrupt = await self.ready(
            await self.upload(b"\x89PNG\r\n\x1a\ncorrupt", "bad.png")
        )
        self.assertEqual(corrupt["processing"]["state"], "failed")
        self.assertTrue(self.service.attachments.read(corrupt["id"])[0].is_file())
        retry = self.service.attachments.retry(corrupt["id"])
        self.assertNotEqual(
            retry["processing"]["preparationId"], corrupt["processing"]["preparationId"]
        )

    async def test_binary_text_and_unsafe_zip_rejected(self):
        for data, name in [
            (b"abc\x00def", "bad.txt"),
            (b"\xff", "bad.txt"),
            (b"old-doc", "old.doc"),
        ]:
            with self.assertRaises(RuntimeFault):
                await self.upload(data, name)
        output = io.BytesIO()
        with ZipFile(output, "w") as archive:
            archive.writestr("[Content_Types].xml", "content")
            archive.writestr("word/document.xml", "content")
            archive.writestr("../evil", "content")
        with self.assertRaises(RuntimeFault):
            await self.upload(output.getvalue(), "bad.docx")

    async def test_claim_is_exclusive_and_late_commit_cannot_overwrite_cancel(self):
        item = await self.ready(await self.upload(b"content", "a.txt"))
        repository = PreparationRepository(self.service.store.database)
        preparation = repository.latest(item["id"])
        self.assertFalse(
            repository.claim(item["id"], preparation["preparationId"], "other")
        )
        self.assertFalse(
            repository.commit(
                preparation["preparationId"], "other", {"state": "failed"}
            )
        )
        self.assertEqual(repository.latest(item["id"])["state"], "ready")

    async def test_image_studio_rejects_documents_before_job_creation(self):
        item = await self.ready(await self.upload(b"text", "a.txt"))
        with self.assertRaises(RuntimeFault) as error:
            self.service.images.attachment_paths([item["id"]])
        self.assertEqual(error.exception.info.code, "IMAGE_FORMAT_UNSUPPORTED")

    async def test_history_visuals_require_validation_honour_none_and_reject_foreign_ids(
        self,
    ):
        output = io.BytesIO()
        Image.new("RGB", (2, 2), "red").save(output, "PNG")
        image = await self.ready(await self.upload(output.getvalue(), "image.png"))
        foreign = await self.ready(await self.upload(output.getvalue(), "foreign.png"))
        conversation = await self.service.chat.create_conversation(
            "Test", "builtin-fake-text"
        )
        self.service.store.conversations.save_message(
            {
                "id": "turn",
                "turnId": "turn",
                "conversationId": conversation["id"],
                "role": "user",
                "content": "image",
                "attachmentIds": [image["id"]],
                "parameters": {
                    "attachmentPreparations": {
                        image["id"]: image["processing"]["preparationId"]
                    }
                },
            }
        )
        conversation = self.service.queries.conversation(conversation["id"])
        model = replace(
            fake_text_manifest(),
            engine="mlx-vlm",
            input_modalities=("text", "image"),
            capabilities=ModelCapabilities(multi_image=True),
        )
        preparation = self.service.chat.preparation
        with self.assertRaises(RuntimeFault) as error:
            preparation.accept(conversation, [], {"mode": "auto"}, model)
        self.assertEqual(error.exception.info.code, "HISTORY_VISUALS_UNVERIFIED")
        verified = replace(
            model, capabilities=replace(model.capabilities, history_images=True)
        )
        accepted = preparation.accept(conversation, [], {"mode": "auto"}, verified)
        self.assertEqual([item["id"] for item in accepted["visuals"]], [image["id"]])
        self.assertEqual(
            preparation.accept(conversation, [], {"mode": "none"}, verified)["visuals"],
            [],
        )
        with self.assertRaises(RuntimeFault):
            preparation.accept(
                conversation,
                [],
                {"mode": "explicit", "attachment_ids": [foreign["id"]]},
                verified,
            )

    async def test_hard_timeout_stops_child_and_can_retry_without_upload(self):
        self.service.attachments.executor.timeout_seconds = 0.001
        item = await self.ready(await self.upload(b"content", "a.txt"))
        self.assertEqual(item["processing"]["errorCode"], "ATTACHMENT_PARSE_TIMEOUT")
        self.service.attachments.executor.timeout_seconds = 30
        retried = await self.ready(self.service.attachments.retry(item["id"]))
        self.assertEqual(retried["id"], item["id"])
        self.assertEqual(retried["processing"]["state"], "ready")
