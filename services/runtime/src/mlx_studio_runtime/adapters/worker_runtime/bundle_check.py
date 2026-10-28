"""Offline, model-free dependency checks executed only by the frozen release sidecar."""

import ctypes
import json
import os
import ssl
import sys
import traceback
from collections.abc import Callable, Generator
from contextlib import contextmanager, redirect_stdout
from functools import partial
from importlib import import_module
from importlib.metadata import version
from io import BytesIO
from pathlib import Path


def check_environment() -> dict[str, object]:
    if not getattr(sys, "frozen", False):
        raise RuntimeError("Bundle verification requires the frozen release sidecar")
    root = getattr(sys, "_MEIPASS", None)
    if not isinstance(root, str) or not Path(root).is_dir():
        raise RuntimeError("Frozen bundle extraction directory is unavailable")
    return {"platform": sys.platform, "python": sys.version.split()[0]}


def check_mlx_device() -> dict[str, str]:
    mx = import_module("mlx.core")
    # Engine imports create streams. This diagnostic must also work on GPU-less CI.
    mx.set_default_device(mx.cpu)
    return {"device": "cpu"}


def check_metadata() -> dict[str, str]:
    versions = {
        package: version(package)
        for package in ("pypdf", "python-docx", "pillow", "mlx-audio", "keyring")
    }
    if versions["mlx-audio"] != "0.5.9":
        raise RuntimeError("Packaged mlx-audio does not match the Qwen3 adapter")
    return versions


def check_imports(names: tuple[str, ...]) -> list[str]:
    for name in names:
        import_module(name)
    return list(names)


def check_documents() -> dict[str, int]:
    from docx import Document
    from pypdf import PdfReader, PdfWriter

    document = Document()  # Reads the packaged default.docx template.
    document.add_paragraph("Ownloom bundle check")
    with BytesIO() as output:
        document.save(output)
        docx_bytes = output.tell()
        output.seek(0)
        if Document(output).paragraphs[0].text != "Ownloom bundle check":
            raise RuntimeError("DOCX resource roundtrip failed")
    with BytesIO() as output, PdfWriter() as writer:
        writer.add_blank_page(width=16, height=16)
        writer.write(output)
        pdf_bytes = output.tell()
        output.seek(0)
        if len(PdfReader(output).pages) != 1:
            raise RuntimeError("PDF resource roundtrip failed")
    return {"docxBytes": docx_bytes, "pdfBytes": pdf_bytes}


def check_images() -> dict[str, int]:
    from PIL import Image

    sizes: dict[str, int] = {}
    for format_name in ("PNG", "JPEG", "GIF", "WEBP"):
        with Image.new("RGB", (2, 2), (12, 34, 56)) as image, BytesIO() as output:
            image.save(output, format_name)
            sizes[format_name] = output.tell()
            output.seek(0)
            with Image.open(output) as decoded:
                decoded.load()
                if decoded.size != (2, 2):
                    raise RuntimeError(f"{format_name} resource roundtrip failed")
    return sizes


def check_tokenizers() -> dict[str, object]:
    tokenizers = import_module("tokenizers")
    models = import_module("tokenizers.models")
    tokenizer = tokenizers.Tokenizer(
        models.WordLevel({"[UNK]": 0, "ownloom": 1}, unk_token="[UNK]")
    )
    if tokenizer.encode("ownloom").ids != [1]:
        raise RuntimeError("Native tokenizer roundtrip failed")
    import_module("tiktoken_ext.openai_public")
    tiktoken = import_module("tiktoken")
    names: list[str] = tiktoken.list_encoding_names()
    if "cl100k_base" not in names:
        raise RuntimeError("Packaged tiktoken encoding plugin is unavailable")
    # get_encoding() may download vocabulary data; discovery itself stays offline.
    transformers = import_module("transformers")
    for name in ("AutoConfig", "AutoTokenizer", "AutoProcessor"):
        if not callable(getattr(transformers, name)):
            raise RuntimeError(f"Transformers lazy export is unavailable: {name}")
    return {"nativeTokenizer": "passed", "tiktokenEncodings": names}


def check_signal() -> dict[str, int]:
    import numpy as np

    signal = import_module("scipy.signal")
    samples = np.array([0.0, 0.5, -0.5, 0.0], dtype=np.float32)
    result = signal.resample_poly(samples, 1, 2)
    if result.shape != (2,) or not np.isfinite(result).all():
        raise RuntimeError("SciPy signal resampling failed")
    return {"inputSamples": len(samples), "outputSamples": len(result)}


def check_resources() -> dict[str, int]:
    root = Path(str(getattr(sys, "_MEIPASS", "")))
    sizes: dict[str, int] = {}
    for name in (
        "mlx/lib/libmlx.dylib",
        "mlx/lib/mlx.metallib",
        "mlx/include/mlx/backend/metal/kernels/steel/gemm/gemm.h",
    ):
        resource = root / name
        if not resource.is_file() or resource.stat().st_size == 0:
            raise RuntimeError(f"Packaged MLX resource is unavailable: {name}")
        sizes[name] = resource.stat().st_size
    import certifi

    context = ssl.create_default_context(cafile=certifi.where())
    if context.cert_store_stats()["x509_ca"] == 0:
        raise RuntimeError("Packaged CA certificates are unavailable")
    return sizes


CHECKS: tuple[tuple[str, Callable[[], object]], ...] = (
    ("mlxDevice", check_mlx_device),
    ("metadata", check_metadata),
    ("documents", check_documents),
    ("images", check_images),
    ("tokenizers", check_tokenizers),
    ("signal", check_signal),
    ("resources", check_resources),
    ("text", partial(check_imports, ("mlx.core", "mlx_lm", "mlx_lm.sample_utils"))),
    (
        "vision",
        partial(
            check_imports,
            ("mlx_vlm", "mlx_vlm.utils", "mlx_vlm.prompt_utils", "mlx_vlm.tokenizer_utils"),
        ),
    ),
    (
        "imageEngines",
        partial(
            check_imports,
            (
                "mflux.models.common.config",
                "mflux.models.flux.variants.txt2img.flux",
                "mflux.models.qwen.variants.txt2img.qwen_image",
                "mflux.models.qwen21.variants.txt2img.qwen_image_21",
                "mflux.models.qwen21.variants.edit.qwen_image_21_edit",
            ),
        ),
    ),
    (
        "speech",
        partial(check_imports, ("mlx_audio.tts.utils", "mlx_audio.tts.models.qwen3_tts")),
    ),
    (
        "nativeSupport",
        partial(
            check_imports,
            ("torch", "cv2", "miniaudio", "sounddevice", "keyring.backends.macOS"),
        ),
    ),
)


@contextmanager
def import_diagnostics() -> Generator[None]:
    """Keep both Python prints and native-library writes off the JSON output pipe."""
    sys.stdout.flush()
    output = sys.stdout.fileno()
    saved = os.dup(output)
    try:
        os.dup2(sys.stderr.fileno(), output)
        with redirect_stdout(sys.stderr):
            yield
    finally:
        try:
            sys.stderr.flush()
            # C stdio may buffer import messages until process exit even after dup2().
            ctypes.CDLL(None).fflush(None)
        finally:
            os.dup2(saved, output)
            os.close(saved)


def run_check(
    name: str, check: Callable[[], object], results: dict[str, dict[str, object]]
) -> None:
    try:
        results[name] = {"status": "passed", "details": check()}
    except Exception as error:
        # This one-shot diagnostic process reports every independent missing dependency.
        results[name] = {"status": "failed", "error": f"{type(error).__name__}: {error}"}
        print(f"Bundle check failed: {name}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)


def main() -> int:
    os.environ.update(HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1", HF_HUB_DISABLE_TELEMETRY="1")
    results: dict[str, dict[str, object]] = {}
    with import_diagnostics():
        run_check("environment", check_environment, results)
        if results["environment"]["status"] == "passed":
            for name, check in CHECKS:
                run_check(name, check, results)
    passed = all(result["status"] == "passed" for result in results.values())
    print(
        json.dumps(
            {
                "profile": "release",
                "scope": "imports-and-resources",
                "device": "cpu",
                "frozen": bool(getattr(sys, "frozen", False)),
                "status": "passed" if passed else "failed",
                "checks": results,
                "unverified": ["model loading and inference", "Metal GPU execution"],
            },
            separators=(",", ":"),
        )
    )
    return 0 if passed else 1
